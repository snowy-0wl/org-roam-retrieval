"""Bounded local extraction. Corpus files are never changed."""
from __future__ import annotations

from functools import lru_cache
import hashlib
import os
import shutil
import json
from pathlib import Path
import re
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET

from .policy import exclusion_reason, extraction_policy
from .formats import html_document, epub_document, latex_document, bibtex_document

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}


@lru_cache(maxsize=8)
def ocr_tools(data_dir=None):
    binary = shutil.which("tesseract")
    if not binary:
        return None, []
    try:
        process = subprocess.run([binary, "--list-langs"] + (["--tessdata-dir", data_dir] if data_dir else []), capture_output=True, text=True, timeout=10, check=True)
        available = set(process.stdout.splitlines()[1:])
        languages = [language for language in ("eng", "heb", "grc", "ell") if language in available]
        return binary, languages
    except (OSError, subprocess.SubprocessError):
        return None, []


def image_ocr(path: str, options=None) -> tuple[str, dict]:
    options = options or {}
    if options.get("enabled", True) is False:
        return "", {"ocr": "disabled", "visual_embeddings": False}
    data_dir = options.get("data_dir")
    binary, available = ocr_tools(data_dir)
    requested = options.get("languages", ["eng", "heb", "grc"])
    languages = [language for language in requested if language in available]
    missing = [language for language in requested if language not in available]
    if not binary or not languages:
        return "", {"ocr": "unavailable", "ocr_languages": languages, "missing_ocr_languages": missing, "visual_embeddings": False}
    environment = dict(os.environ, OMP_THREAD_LIMIT="1")
    try:
        process = subprocess.run([binary, path, "stdout", "-l", "+".join(languages)] + (["--tessdata-dir", data_dir] if data_dir else []),
                                 capture_output=True, timeout=45, check=True, env=environment)
        return process.stdout.decode("utf-8", errors="replace"), {"ocr": "text", "ocr_languages": languages, "missing_ocr_languages": missing, "visual_embeddings": False}
    except (OSError, subprocess.SubprocessError) as error:
        return "", {"ocr": "failed", "ocr_languages": languages, "missing_ocr_languages": missing, "visual_embeddings": False,
                    "ocr_error": type(error).__name__}


def pdf_page_ocr(path: str, page: int, options=None) -> tuple[str, dict]:
    if options and options.get("enabled", True) is False:
        return "", {"ocr": "disabled", "visual_embeddings": False}
    renderer = shutil.which("pdftoppm")
    if not renderer:
        return "", {"ocr": "renderer-unavailable", "visual_embeddings": False}
    with tempfile.TemporaryDirectory(prefix="org-roam-page-") as temporary:
        prefix = str(Path(temporary) / "page")
        try:
            subprocess.run([renderer, "-f", str(page), "-l", str(page), "-scale-to", "2000", "-singlefile", "-png", path, prefix],
                           capture_output=True, timeout=60, check=True)
            return image_ocr(prefix + ".png", options)
        except (OSError, subprocess.SubprocessError) as error:
            return "", {"ocr": "render-failed", "ocr_error": type(error).__name__, "visual_embeddings": False}


TEXT_EXTENSIONS = {".txt", ".text", ".md", ".markdown", ".rst", ".csv", ".tsv"}


def chunks(text: str, *, size: int = 1600, overlap: int = 160):
    """Cover every character, using paragraph/word boundaries where possible."""
    text = text.strip()
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = max(text.rfind("\n", start + size // 2, end), text.rfind(" ", start + size // 2, end))
            if boundary > start:
                end = boundary
        yield text[start:end]
        if end == len(text):
            break
        start = max(start + 1, end - overlap)


def org_batch(paths: list[str], emacs: str, extractor: str) -> tuple[list[dict], list[dict]]:
    if not paths:
        return [], []
    with tempfile.TemporaryDirectory(prefix="org-roam-extract-") as temporary:
        request = Path(temporary) / "request.json"
        request.write_text(json.dumps({"files": paths}, ensure_ascii=False), encoding="utf-8")
        request.chmod(0o600)
        expression = f"(org-roam-retrieval-extract-batch {json.dumps(str(request))})"
        completed = subprocess.run(
            [emacs, "-Q", "--batch", "-L", str(Path(extractor).parent), "-l", extractor, "--eval", expression],
            capture_output=True, timeout=120, check=False,
        )
        if completed.returncode:
            # Only error-category lines are retained; backtraces can contain source text.
            lines = completed.stderr.decode("utf-8", errors="replace").splitlines()
            category = next((line for line in reversed(lines) if line.startswith(("Error:", "Cannot open", "Symbol’s", "Invalid read syntax"))), "No safe diagnostic; inspect extractor validation")
            raise RuntimeError(f"Org helper exited {completed.returncode}: {category[:300]}")
        result = json.loads(completed.stdout.decode("utf-8"))
        return result.get("sources", []), result.get("errors", [])


def document(path: str, ocr=None) -> dict:
    file = Path(path)
    suffix = file.suffix.lower()
    title = file.stem
    if reason := exclusion_reason(path):
        return {"path": path, "title": title, "segments": [], "coverage": "excluded",
                "details": {"exclusion_reason": reason, "extraction_policy": extraction_policy(path)}}
    if suffix in {".html", ".htm", ".xhtml"}:
        return html_document(path)
    if suffix == ".epub":
        return epub_document(path)
    if suffix in {".tex", ".latex"}:
        return latex_document(path)
    if suffix == ".bib":
        return bibtex_document(path)
    segments: list[dict] = []
    coverage = "text"
    if suffix in TEXT_EXTENSIONS:
        text = file.read_text(encoding="utf-8", errors="replace")
        if suffix in {".md", ".markdown"}:
            text = re.sub(r"(?ms)^```.*?^```[^\n]*", "", text)
            text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
        segments.append({"text": text, "title": title, "locator": {"kind": "text-range", "line": 1}})
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(file)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("Password-protected PDF requires a password; source left unchanged")
        empty_pages = []
        ocr_pages = []
        ocr_details = {}
        for number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                text, ocr_details = pdf_page_ocr(path, number, ocr)
                if text.strip():
                    ocr_pages.append(number)
            if text.strip():
                segments.append({"text": text, "title": title, "locator": {"kind": "pdf-page", "page": number}})
            else:
                empty_pages.append(number)
        if empty_pages:
            coverage = "partial-text" if segments else "filename-only"
        return {"path": path, "title": title, "segments": segments, "coverage": coverage,
                "details": {"pages": len(reader.pages), "empty_pages": empty_pages, "ocr_pages": ocr_pages, **ocr_details}}
    elif suffix in {".docx", ".odt"}:
        member = "word/document.xml" if suffix == ".docx" else "content.xml"
        with zipfile.ZipFile(file) as archive:
            if archive.getinfo(member).file_size > 64 * 1024 * 1024:
                raise ValueError("Office document XML exceeds the 64 MiB extraction limit")
            xml = ET.fromstring(archive.read(member))
        if suffix == ".docx":
            paragraphs = ["".join(node.itertext()) for node in xml.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")]
        else:
            paragraphs = ["".join(node.itertext()) for node in xml.iter() if node.tag in
                          ("{urn:oasis:names:tc:opendocument:xmlns:text:1.0}p", "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}h")]
        segments.append({"text": "\n".join(paragraphs), "title": title, "locator": {"kind": "file"}})
    elif suffix in IMAGE_EXTENSIONS:
        text, details = image_ocr(path, ocr)
        if text.strip():
            segments.append({"text": text, "title": title, "locator": {"kind": "file"}})
            coverage = "ocr-text"
        else:
            coverage = "filename-only"
            segments.append({"text": title, "title": title, "locator": {"kind": "file"}})
        return {"path": path, "title": title, "segments": segments, "coverage": coverage, "details": details}
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(file, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook:
                rows = []
                first = 1
                for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = [str(value) if value is not None else "" for value in row]
                    if any(values):
                        rows.append("\t".join(values))
                    if len(rows) >= 30:
                        segments.append({"text": "\n".join(rows), "title": f"{title} · {sheet.title}",
                                         "locator": {"kind": "sheet-range", "sheet": sheet.title, "range": f"A{first}:XFD{number}"}})
                        first, rows = number + 1, []
                if rows:
                    segments.append({"text": "\n".join(rows), "title": f"{title} · {sheet.title}",
                                     "locator": {"kind": "sheet-range", "sheet": sheet.title, "range": f"A{first}:XFD{sheet.max_row}"}})
        finally:
            workbook.close()
    else:
        coverage = "filename-only"
    details = {"extraction_status": "unsupported-format", "unsupported_format": suffix or "extensionless"} if coverage == "filename-only" else {}
    if not segments:
        segments.append({"text": title, "title": title, "locator": {"kind": "file"}})
    return {"path": path, "title": title, "segments": segments, "coverage": coverage, "details": details}


def source_id(path: str) -> str:
    return hashlib.sha256(str(Path(path).resolve()).encode()).hexdigest()


def prepare_chunks(source: dict) -> list[dict]:
    prepared = []
    for segment_number, segment in enumerate(source.get("segments", [])):
        fulltext = segment.get("text", "")
        last_offset = 0
        for number, text in enumerate(chunks(fulltext)):
            locator = dict(segment["locator"])
            offset = fulltext.find(text, last_offset)
            if offset < 0:
                offset = last_offset
            last_offset = offset + 1
            if locator["kind"] == "text-range":
                locator["line"] = int(locator.get("line", 1)) + fulltext[:offset].count("\n")
            title = segment.get("title") or source["title"]
            # The bounded title/body envelope stays well below the 8192-token model cap;
            # the actual tokenizer is checked again before every inference.
            embedding_input = f"title: {title[:200]} | text: {text}"
            prepared.append({"title": title, "text": text, "locator": locator,
                             "input": embedding_input,
                             "input_hash": hashlib.sha256(embedding_input.encode()).hexdigest(),
                             "ordinal": f"{segment_number}:{number}"})
    return prepared
